# -*- coding: utf-8 -*-
"""从《焦圈儿模板数据_prmpt.xlsx》导入提示词与模板库。

用法（在 backend/ 目录下）：
    python -m scripts.import_prompts --dry    # 干跑：只打印合并计划，不写库
    python -m scripts.import_prompts          # 实际导入（需先跑过 python -m app.seeds）

合并策略：
- 按模板名称匹配现有模板：存在 → 挂载新 prompt_version 并发布（内容相同则跳过，幂等）
- 不存在 → 新建模板（元数据取自 xlsx：描述/评分/焦耳价格/使用示例/分类挂载/推荐排序）
- 重复执行安全：版本内容不变则不新增版本
"""
import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]          # MyJQR/
XLSX = ROOT / "docs" / "焦圈儿模板数据_prmpt.xlsx"

USER_PROMPT_TMPL = "主题：{topic}\n素材：{material}\n字数：约{length}字\n请按系统设定创作。"


def read_xlsx(xlsx_path: Path) -> dict[str, dict]:
    """返回 {模板名: {prompt, description, score, categories[], rank, price, example, tips}}"""
    wb = load_workbook(xlsx_path, read_only=True)

    # ① 分类模板表：提示词 + 分类挂载
    ws = wb["分类模板"]
    out: dict[str, dict] = {}
    for r in range(5, ws.max_row + 1):
        cat = ws.cell(row=r, column=2).value    # 分类名称
        name = ws.cell(row=r, column=6).value   # 模板名称
        score = ws.cell(row=r, column=7).value  # 评分
        desc = ws.cell(row=r, column=8).value   # 描述
        prompt = ws.cell(row=r, column=16).value  # 提示词
        if not name:
            continue
        name = str(name).strip()
        item = out.setdefault(name, {"prompt": "", "description": "", "score": 5.0,
                                     "categories": [], "rank": 99, "price": None,
                                     "example": "", "tips": ""})
        if cat and str(cat).strip() not in item["categories"]:
            item["categories"].append(str(cat).strip())
        if prompt and str(prompt).strip():
            item["prompt"] = str(prompt).strip()
        if desc and not item["description"]:
            item["description"] = str(desc).strip()
        if isinstance(score, (int, float)):
            item["score"] = max(0, min(score, 5))

    # ② 推荐模板表：推荐排序
    ws = wb["推荐模板"]
    for r in range(5, ws.max_row + 1):
        rank = ws.cell(row=r, column=2).value
        name = ws.cell(row=r, column=4).value
        if name and isinstance(rank, int) and str(name).strip() in out:
            out[str(name).strip()]["rank"] = min(out[str(name).strip()]["rank"], rank)

    # ③ 模板详情表：焦耳价格 / 使用示例 / 使用提示
    ws = wb["模板详情"]
    for r in range(5, ws.max_row + 1):
        name = ws.cell(row=r, column=3).value
        if not name or str(name).strip() not in out:
            continue
        item = out[str(name).strip()]
        price = ws.cell(row=r, column=10).value
        if isinstance(price, (int, float)) and item["price"] is None:
            item["price"] = int(price)
        ex, tips = ws.cell(row=r, column=7).value, ws.cell(row=r, column=8).value
        if ex and not item["example"]:
            item["example"] = str(ex).strip()
        if tips and not item["tips"]:
            item["tips"] = str(tips).strip()
    return out


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry", action="store_true", help="只打印计划，不写库")
    parser.add_argument("--xlsx", default=str(XLSX), help="xlsx 路径")
    args = parser.parse_args()

    items = read_xlsx(Path(args.xlsx))
    with_prompt = {k: v for k, v in items.items() if v["prompt"]}
    print(f"[parse] 唯一模板 {len(items)} 个，其中含提示词 {len(with_prompt)} 个")
    no_prompt = sorted(set(items) - set(with_prompt))
    if no_prompt:
        print(f"[parse] 无提示词（跳过或沿用旧版）：{'、'.join(no_prompt)}")

    if args.dry:
        for name, v in sorted(with_prompt.items(), key=lambda x: x[1]["rank"]):
            print(f"  - {name:<14} 分类:{'/'.join(v['categories']) or '—':<10} "
                  f"价格:{v['price'] or '—':>5}J 排名:{v['rank']:>2} "
                  f"提示词{len(v['prompt'])}字")
        print("[dry] 干跑结束，未写库。去掉 --dry 执行导入。")
        return

    from app.core.config import get_settings
    from app.db import SessionLocal, engine, Base
    from app.models.template import (Category, PromptVersion, Template,
                                     TemplateCategoryRelation, TemplateStat)
    Base.metadata.create_all(engine)
    db = SessionLocal()
    try:
        cats = {c.name: c.id for c in db.query(Category).all()}
        created = updated = skipped = 0
        for name, v in sorted(with_prompt.items(), key=lambda x: x[1]["rank"]):
            tpl = db.query(Template).filter_by(name=name).first()
            if not tpl:
                tpl = Template(name=name, description=v["description"],
                               avatar_url="", author_id=0, workflow_type="single",
                               model_route="default_chat",
                               enable_web_search=1, enable_think=1,
                               temperature=0.95, max_tokens=8192,
                               response_length=2, joule_price=v["price"] or 500,
                               score=v["score"], status=1, sort_order=v["rank"],
                               usage_example=v["example"], usage_tips=v["tips"])
                db.add(tpl)
                db.flush()
                db.add(TemplateStat(template_id=tpl.id))
                for r_, cname in enumerate(v["categories"]):
                    if cname in cats:
                        db.add(TemplateCategoryRelation(template_id=tpl.id,
                                                        category_id=cats[cname], rank=r_))
                created += 1
            else:
                tpl.description = tpl.description or v["description"]
                tpl.usage_example = tpl.usage_example or v["example"]
                tpl.usage_tips = tpl.usage_tips or v["tips"]
                if v["price"]:
                    tpl.joule_price = v["price"]
                tpl.score = v["score"]
                tpl.sort_order = v["rank"]

            latest = (db.query(PromptVersion).filter_by(template_id=tpl.id)
                      .order_by(PromptVersion.version.desc()).first())
            if latest and latest.system_prompt == v["prompt"]:
                if not latest.published:
                    latest.published = 1
                skipped += 1
                continue
            if latest:
                db.query(PromptVersion).filter_by(template_id=tpl.id).update({"published": 0})
                version = latest.version + 1
            else:
                version = 1
            db.add(PromptVersion(template_id=tpl.id, version=version, published=1,
                                 system_prompt=v["prompt"],
                                 user_prompt_template=USER_PROMPT_TMPL))
            updated += 1
        db.commit()
        print(f"[import] 新建模板 {created}，更新提示词 {updated}，内容未变跳过 {skipped}")
        print(f"[import] 完成：库内模板总数 {db.query(Template).count()}")
    finally:
        db.close()


if __name__ == "__main__":
    main()

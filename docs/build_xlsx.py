# -*- coding: utf-8 -*-
"""焦圈儿模板数据 → Excel（4 工作表：分类 / 推荐模板 / 分类模板 / 模板详情）"""
import sys, os, json, datetime

XLSX_SKILL_DIR = r"C:\Users\飞象\.zcode\cli\plugins\cache\zcode-plugins-official\document-skills\0.1.4\skills\xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
from base import (FONT_NAME, setup_sheet, style_header_row, style_data_row, style_total_row,
                  auto_fit_columns, auto_fit_row_heights, font_caption, align_number)
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def enforce_widths(ws, widths):
    """在 auto_fit 之后按需加宽指定列（col_letter -> 最小宽度）"""
    for letter, w in widths.items():
        cur = ws.column_dimensions[letter].width or 0
        if cur < w:
            ws.column_dimensions[letter].width = w

with open("raw_data.json", encoding="utf-8") as f:
    RAW = json.load(f)

cats = RAW["categories"]
rec = RAW["recommend"]
cat_tpl = RAW["category_templates"]
details = RAW["details"]

cat_name = {c["category_id"]: c["name"] for c in cats}
rec_rank = {t["template_id"]: i + 1 for i, t in enumerate(rec)}
tpl_cats = {}
for cid, lst in cat_tpl.items():
    for t in lst:
        tpl_cats.setdefault(t["template_id"], []).append(cat_name[cid])


def s(v):
    return (str(v) if v is not None else "").strip()


def yn(v):
    return "是" if v else "否"


def hot_titles(t, limit=3):
    items = t.get("hot_contents") or []
    return "\n".join(f"{i+1}. {s(h.get('title'))}" for i, h in enumerate(items[:limit]))


def tags_join(t):
    tg = t.get("tags") or []
    return "、".join(s(x.get("name") if isinstance(x, dict) else x) for x in tg)


def cfg(t, k):
    return s((t.get("recommend_config") or {}).get(k))


def author(t, k):
    return s((t.get("author") or {}).get(k))


def stats(t, k):
    v = (t.get("stats") or {}).get(k)
    return v if isinstance(v, (int, float)) else None


SOURCE_NOTE = f"数据来源：api-server.jiaoquaner.com（studio.jiaoquaner.com）｜抓取时间：{datetime.date.today().isoformat()}｜分页方式：cursor"

wb = Workbook()


def build_sheet(ws, title, headers, rows, num_cols=None, num_fmt=None, note=True, freeze="B5"):
    """headers 从 B4 起，数据从 B5 起；num_cols=需要右对齐的列号(绝对列)集合"""
    last_col = len(headers) + 1
    setup_sheet(ws, title=title, last_col=last_col)
    for ci, h in enumerate(headers, start=2):
        ws.cell(row=4, column=ci, value=h)
    style_header_row(ws, row_num=4, col_start=2, col_end=last_col)
    for i, row in enumerate(rows):
        rn = 5 + i
        for ci, v in enumerate(row, start=2):
            ws.cell(row=rn, column=ci, value=v)
        style_data_row(ws, row_num=rn, col_start=2, col_end=last_col, row_index=i)
        if num_cols:
            for ci in num_cols:
                c = ws.cell(row=rn, column=ci)
                c.alignment = align_number()
                if num_fmt and ci in num_fmt:
                    c.number_format = num_fmt[ci]
    end_row = 4 + len(rows)
    if note and rows:
        nc = ws.cell(row=end_row + 2, column=2, value=SOURCE_NOTE)
        nc.font = font_caption()
    ws.freeze_panes = freeze
    auto_fit_columns(ws, min_width=8, max_width=38, header_row=4, data_start_row=5)
    auto_fit_row_heights(ws, header_row=4, data_start_row=5)
    return end_row


# ── Sheet 1: 分类 ──────────────────────────────────────────────
ws1 = wb.active
ws1.title = "分类"
rows1 = [[int(c["category_id"]), c["name"], int(c["sort_order"]), len(cat_tpl.get(c["category_id"], []))]
         for c in sorted(cats, key=lambda x: int(x["sort_order"]))]
h1 = ["分类ID", "分类名称", "排序", "模板数量"]
end1 = build_sheet(ws1, "焦圈儿模板分类一览", h1, rows1, num_cols={2, 4, 5},
                   num_fmt={5: "#,##0"}, note=False)
enforce_widths(ws1, {"B": 9, "C": 11, "D": 8, "E": 10})
tr = end1 + 1
ws1.cell(row=tr, column=2, value="合计")
ws1.cell(row=tr, column=5, value=f"=SUM(E5:E{end1})")
style_total_row(ws1, row_num=tr, col_start=2, col_end=5)
ws1.cell(row=tr, column=5).alignment = align_number()
ws1.cell(row=tr, column=5).number_format = "#,##0"
cap = ws1.cell(row=tr + 2, column=2, value=SOURCE_NOTE + "｜合计为分类-模板关系数（同一模板可属多个分类，去重后模板共 55 个）")
cap.font = font_caption()

# ── Sheet 2: 推荐模板 ─────────────────────────────────────────
ws2 = wb.create_sheet("推荐模板")
rows2 = []
for i, t in enumerate(rec):
    rows2.append([
        i + 1, int(t["template_id"]), t["name"], t.get("score"),
        s(t.get("description")),
        cfg(t, "model_display_name") or cfg(t, "model"),
        yn(cfg(t, "enable_web_search") == "True"), yn(cfg(t, "enable_think") == "True"),
        tags_join(t), author(t, "nickname"), author(t, "user_id") or "—",
        stats(t, "user_count"), stats(t, "use_count"), s((t.get("stats") or {}).get("use_count_text")),
        hot_titles(t), yn(t.get("is_bookmarked")),
        "、".join(tpl_cats.get(t["template_id"], [])),
    ])
h2 = ["推荐排名", "模板ID", "模板名称", "评分", "描述", "模型", "联网搜索", "深度思考", "标签",
      "作者昵称", "作者ID", "用户数", "使用次数", "使用次数文案", "热门内容TOP3", "已收藏", "所属分类"]
build_sheet(ws2, "推荐模板（recommend/list/page，按推荐顺序）", h2, rows2,
            num_cols={2, 5, 13, 14}, num_fmt={5: "0.0", 13: "#,##0", 14: "#,##0"})
enforce_widths(ws2, {"E": 8, "G": 11, "L": 22, "M": 11, "N": 11})

# ── Sheet 3: 分类模板 ─────────────────────────────────────────
ws3 = wb.create_sheet("分类模板")
rows3 = []
for c in sorted(cats, key=lambda x: int(x["sort_order"])):
    for j, t in enumerate(cat_tpl.get(c["category_id"], [])):
        rows3.append([
            c["name"], int(c["category_id"]), j + 1, int(t["template_id"]), t["name"],
            t.get("score"), s(t.get("description")),
            cfg(t, "model_display_name") or cfg(t, "model"),
            author(t, "nickname"), tags_join(t),
            stats(t, "user_count"), stats(t, "use_count"), hot_titles(t),
            f"是(第{rec_rank[t['template_id']]}位)" if t["template_id"] in rec_rank else "否",
        ])
h3 = ["分类名称", "分类ID", "分类内排名", "模板ID", "模板名称", "评分", "描述", "模型",
      "作者昵称", "标签", "用户数", "使用次数", "热门内容TOP3", "是否推荐"]
build_sheet(ws3, "各分类模板（template/list?category_id=…）", h3, rows3,
            num_cols={3, 4, 6, 12, 13}, num_fmt={6: "0.0", 12: "#,##0", 13: "#,##0"})
enforce_widths(ws3, {"C": 9, "D": 10, "I": 11, "L": 11, "M": 11})

# ── Sheet 4: 模板详情 ─────────────────────────────────────────
ws4 = wb.create_sheet("模板详情")
rows4 = []
for tid in sorted(details, key=lambda x: int(x)):
    d = details[tid]
    rows4.append([
        int(tid), d["name"],
        f"是(第{rec_rank[tid]}位)" if tid in rec_rank else "否",
        "、".join(tpl_cats.get(tid, [])),
        s(d.get("description")),
        cfg(d, "model_display_name") or cfg(d, "model"),
        s(d.get("usage_example")), s(d.get("usage_tips")),
        d.get("joule_price") if isinstance(d.get("joule_price"), (int, float)) else None,
        hot_titles(d),
    ])
h4 = ["模板ID", "模板名称", "是否推荐", "所属分类", "描述", "模型", "使用示例", "使用提示",
      "焦耳价格", "热门内容TOP3"]
build_sheet(ws4, "模板详情（template/detail，全部唯一模板）", h4, rows4,
            num_cols={2, 10}, num_fmt={10: "#,##0"})
enforce_widths(ws4, {"G": 11, "J": 10})

wb.properties.creator = "Z.ai"
out = r"C:\Users\飞象\Desktop\JQE\焦圈儿模板数据.xlsx"
wb.save(out)
print("saved:", out)
print("rows:", {"分类": len(rows1), "推荐模板": len(rows2), "分类模板": len(rows3), "模板详情": len(rows4)})
